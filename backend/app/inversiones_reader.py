"""Module to read Inversiones Excel files and return structured data."""
import openpyxl
from pathlib import Path

DATA_DIR = Path(__file__).parent.parent.parent / "data" / "inversiones"


def get_ahorro() -> list[dict]:
    """Read savings accounts from Excel."""
    filepath = DATA_DIR / "ahorro.xlsx"
    if not filepath.exists():
        return []

    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active

    accounts = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        try:
            name = str(row[0])
            description = str(row[1]) if row[1] else ""
            color = str(row[2]) if row[2] else "#1da1f2"
            balance = float(row[3]) if row[3] else 0
            annual_rate = float(row[4]) if row[4] else 0
            daily_gain = (balance * annual_rate / 100) / 365

            accounts.append({
                "name": name,
                "description": description,
                "color": color,
                "balance": balance,
                "annualRate": annual_rate,
                "dailyGain": round(daily_gain, 4)
            })
        except (ValueError, TypeError, IndexError):
            continue

    wb.close()
    return accounts


def get_prestamos() -> list[dict]:
    """Read loans from Excel and calculate interest and total return.
    
    The interest calculation considers:
    - Rate is annual (%)
    - Term is in months
    - Formula: Interest = Principal × (Annual_Rate / 100) × (Months / 12)
    """
    filepath = DATA_DIR / "prestamos.xlsx"
    if not filepath.exists():
        return []

    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active

    loans = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        if not row or row[0] is None:
            continue
        try:
            principal = float(row[1]) if row[1] else 0
            rate = float(row[2]) if row[2] else 0
            term_str = str(row[3]) if row[3] else ""
            
            # Extraer el número de meses del plazo
            months = _parse_term_to_months(term_str)
            
            # Calcular interés esperado (tasa anual prorrateada por meses)
            # Interés = Capital × (Tasa_Anual / 100) × (Meses / 12)
            expected_interest = principal * (rate / 100) * (months / 12)
            total_return = principal + expected_interest
            
            # Normalizar el estado
            status_raw = str(row[4]).lower().strip() if row[4] else "activo"
            
            # Determinar si el préstamo está activo
            # Estados activos: "activo", "en curso", "vigente"
            # Estados inactivos: "pagado", "cancelado", "finalizado"
            is_active = status_raw in ["activo", "en curso", "vigente"]
            status_label = "Activo" if is_active else "Pagado"
            
            loans.append({
                "id": idx,
                "borrower": str(row[0]),
                "principal": principal,
                "rate": rate,
                "term": term_str,
                "expectedInterest": round(expected_interest, 2),
                "totalReturn": round(total_return, 2),
                "status": "activo" if is_active else "pagado",  # Normalizado para el frontend
                "statusLabel": status_label
            })
        except (ValueError, TypeError, IndexError):
            continue

    wb.close()
    return loans


def _parse_term_to_months(term_str: str) -> int:
    """Parse term string to number of months.
    
    Examples:
        "6 meses" -> 6
        "1 año" -> 12
        "2 años" -> 24
        "18" -> 18
    """
    import re
    
    term_lower = term_str.lower().strip()
    
    # Buscar número en el string
    numbers = re.findall(r'\d+', term_lower)
    if not numbers:
        return 12  # Default a 1 año si no se encuentra número
    
    num = int(numbers[0])
    
    # Si contiene "año" o "year", multiplicar por 12
    if 'año' in term_lower or 'year' in term_lower:
        return num * 12
    
    # Si contiene "mes" o "month", usar el número directamente
    if 'mes' in term_lower or 'month' in term_lower:
        return num
    
    # Si solo es un número sin unidad, asumir que son meses
    return num


def get_afore() -> dict:
    """Read afore data from Excel."""
    filepath = DATA_DIR / "afore.xlsx"
    if not filepath.exists():
        return {"balance": 0, "annualReturn": 0, "bimonthlyContribution": 0, "voluntaryContribution": 0}

    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active

    data = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        field = str(row[0]).strip()
        value = float(row[1]) if row[1] else 0

        if "saldo" in field.lower():
            data["balance"] = value
        elif "rendimiento" in field.lower():
            data["annualReturn"] = value
        elif "bimestral" in field.lower():
            data["bimonthlyContribution"] = value
        elif "voluntaria" in field.lower() or "semanal" in field.lower():
            data["voluntaryContribution"] = value

    wb.close()
    return data


def get_all_inversiones() -> dict:
    """Get all investment data."""
    ahorro = get_ahorro()
    prestamos = get_prestamos()
    afore = get_afore()

    total_savings = sum(a["balance"] for a in ahorro)
    active_loans = [l for l in prestamos if l["status"] == "activo"]
    total_loans = sum(l["principal"] for l in active_loans)
    total_loan_interest = sum(l["expectedInterest"] for l in active_loans)
    avg_loan_rate = (sum(l["rate"] for l in active_loans) / len(active_loans)) if active_loans else 0

    return {
        "ahorro": ahorro,
        "prestamos": prestamos,
        "afore": afore,
        "summary": {
            "totalSavings": total_savings,
            "totalLoans": total_loans,
            "totalLoanInterest": total_loan_interest,
            "avgLoanRate": round(avg_loan_rate, 1)
        }
    }
