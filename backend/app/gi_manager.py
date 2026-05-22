"""Module to manage Gastos/Ingresos Excel file as database."""
import os
from pathlib import Path
from datetime import datetime
import openpyxl

DATA_DIR = Path(__file__).parent.parent.parent / "data" / "GI"
GI_FILE = DATA_DIR / "gastos-ingresos.xlsx"

HEADERS = ["ID", "Fecha", "Descripción", "Categoría", "Tipo", "Monto"]


def ensure_file_exists():
    """Create the Excel file with headers if it doesn't exist."""
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    if not GI_FILE.exists():
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Gastos e Ingresos"
        ws.append(HEADERS)
        wb.save(GI_FILE)


def get_next_id() -> int:
    """Get the next available ID."""
    ensure_file_exists()
    wb = openpyxl.load_workbook(GI_FILE)
    ws = wb.active
    max_id = 0
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        if row[0] and isinstance(row[0], (int, float)):
            max_id = max(max_id, int(row[0]))
    wb.close()
    return max_id + 1


def add_record(fecha: str, descripcion: str, categoria: str, tipo: str, monto: float) -> dict:
    """Add a new record to the Excel file."""
    ensure_file_exists()

    record_id = get_next_id()

    # Monto negativo para gastos, positivo para ingresos
    if tipo == "gasto":
        monto = -abs(monto)
    else:
        monto = abs(monto)

    wb = openpyxl.load_workbook(GI_FILE)
    ws = wb.active
    ws.append([record_id, fecha, descripcion, categoria, tipo, monto])
    wb.save(GI_FILE)
    wb.close()

    return {
        "id": record_id,
        "date": fecha,
        "description": descripcion,
        "category": categoria,
        "type": tipo,
        "amount": monto
    }


def get_all_records() -> list[dict]:
    """Read all records from the Excel file."""
    ensure_file_exists()

    wb = openpyxl.load_workbook(GI_FILE, read_only=True)
    ws = wb.active

    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        try:
            records.append({
                "id": int(row[0]),
                "date": str(row[1]) if row[1] else "",
                "description": str(row[2]) if row[2] else "",
                "category": str(row[3]) if row[3] else "",
                "type": str(row[4]) if row[4] else "",
                "amount": float(row[5]) if row[5] else 0
            })
        except (ValueError, TypeError, IndexError):
            continue

    wb.close()
    # Return sorted by date descending
    records.sort(key=lambda r: r["date"], reverse=True)
    return records


def delete_record(record_id: int) -> bool:
    """Delete a record by ID."""
    ensure_file_exists()

    wb = openpyxl.load_workbook(GI_FILE)
    ws = wb.active

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_col=1, values_only=True), start=2):
        if row[0] and int(row[0]) == record_id:
            ws.delete_rows(row_idx)
            wb.save(GI_FILE)
            wb.close()
            return True

    wb.close()
    return False
