"""Module to read GBM portfolio data from SQLite, with Excel upload support."""
import openpyxl
from io import BytesIO
from app.database import get_db
from app import cache

# Tipo de cambio USD/MXN por defecto (fallback si falla la consulta)
_USD_MXN_FALLBACK = 19.45


def _parse_num(value) -> float:
    """Convierte un valor a float, manejando strings con formato de moneda/porcentaje."""
    if value is None or value == "-":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    if s in ("-", "", "N/A"):
        return 0.0
    negative = False
    if s.startswith("-"):
        negative = True
        s = s[1:]
    s = s.replace("$", "").replace("%", "").replace(",", "").strip()
    try:
        result = float(s)
        return -result if negative else result
    except (ValueError, TypeError):
        return 0.0


def fetch_usd_mxn() -> float:
    """Obtiene el tipo de cambio USD/MXN en tiempo real con caché de 10 minutos."""
    cached = cache.get("usd_mxn_rate")
    if cached is not None:
        return cached

    try:
        import httpx
        url = "https://api.frankfurter.app/latest?base=USD&symbols=MXN"
        resp = httpx.get(url, timeout=5, follow_redirects=True)
        resp.raise_for_status()
        data = resp.json()
        rate = round(float(data["rates"]["MXN"]), 4)
        cache.set("usd_mxn_rate", rate, ttl_seconds=600)
        return rate
    except Exception:
        return _USD_MXN_FALLBACK


def parse_nacional_excel(content: bytes) -> list[dict]:
    """Parse national market portfolio from Excel bytes."""
    wb = openpyxl.load_workbook(BytesIO(content), read_only=True)
    ws = wb.active

    instruments = []
    current_section = ""

    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue

        cell0 = str(row[0]).strip()

        if cell0 in ("Mercado de Capitales Nacional", "Fondos de Inversión Deuda", "Efectivo"):
            current_section = cell0
            continue

        if cell0 in ("Emisora/Fondo", "App GBM Portfolio"):
            continue

        if cell0.startswith("EFEC.") and _parse_num(row[5]) == 0:
            continue

        try:
            ticker = cell0
            shares = _parse_num(row[1])
            avg_cost = _parse_num(row[2])
            market_price = _parse_num(row[3])
            market_value = _parse_num(row[5])
            gain_loss = _parse_num(row[6])
            var_day_pct = _parse_num(row[8])
            portfolio_pct = _parse_num(row[10])

            if avg_cost > 0 and shares > 0:
                total_cost = avg_cost * shares
                return_pct = ((market_value - total_cost) / total_cost) * 100
            else:
                return_pct = 0

            instruments.append({
                "ticker": ticker,
                "section": current_section,
                "shares": shares,
                "avgCost": avg_cost,
                "marketPrice": market_price,
                "marketValue": market_value,
                "gainLoss": gain_loss,
                "returnPct": round(return_pct, 2),
                "varDayPct": var_day_pct,
                "portfolioPct": portfolio_pct
            })
        except (ValueError, TypeError, IndexError):
            continue

    wb.close()
    return instruments


def parse_usa_excel(content: bytes) -> list[dict]:
    """Parse USA market portfolio from Excel bytes."""
    wb = openpyxl.load_workbook(BytesIO(content), read_only=True)
    ws = wb.active

    instruments = []
    current_section = ""

    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue

        cell0 = str(row[0]).strip()

        if cell0 in ("Mercado de Capitales USA", "Liquidez"):
            current_section = cell0
            continue

        if cell0 in ("Emisora/Fondo", "App GBM Portfolio"):
            continue

        if cell0 == "efectivo" and _parse_num(row[5]) == 0:
            continue

        try:
            ticker = cell0
            shares = _parse_num(row[1])
            avg_cost = _parse_num(row[2])
            market_price = _parse_num(row[3])
            market_value = _parse_num(row[5])
            gain_loss = _parse_num(row[6])
            var_day_pct = _parse_num(row[8])
            cost_value = _parse_num(row[9])
            portfolio_pct = _parse_num(row[10])

            instruments.append({
                "ticker": ticker,
                "section": current_section,
                "shares": shares,
                "avgCost": avg_cost,
                "marketPrice": market_price,
                "marketValue": market_value,
                "gainLoss": gain_loss,
                "returnPct": round(((market_price - avg_cost) / avg_cost) * 100, 2) if avg_cost > 0 else 0,
                "varDayPct": var_day_pct,
                "costValue": cost_value,
                "portfolioPct": portfolio_pct
            })
        except (ValueError, TypeError, IndexError):
            continue

    wb.close()
    return instruments


def save_gbm_data(nacional: list[dict], usa: list[dict]):
    """Replace all GBM data in SQLite with new parsed data."""
    with get_db() as conn:
        # Clear old data
        conn.execute("DELETE FROM gbm_nacional")
        conn.execute("DELETE FROM gbm_usa")

        # Insert nacional
        for inv in nacional:
            conn.execute(
                """INSERT INTO gbm_nacional (ticker, section, shares, avg_cost, market_price, market_value, gain_loss, return_pct, var_day_pct, portfolio_pct)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (inv["ticker"], inv["section"], inv["shares"], inv["avgCost"], inv["marketPrice"],
                 inv["marketValue"], inv["gainLoss"], inv["returnPct"], inv["varDayPct"], inv["portfolioPct"])
            )

        # Insert usa
        for inv in usa:
            conn.execute(
                """INSERT INTO gbm_usa (ticker, section, shares, avg_cost, market_price, market_value, gain_loss, return_pct, var_day_pct, cost_value, portfolio_pct)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (inv["ticker"], inv["section"], inv["shares"], inv["avgCost"], inv["marketPrice"],
                 inv["marketValue"], inv["gainLoss"], inv["returnPct"], inv["varDayPct"],
                 inv.get("costValue", 0), inv["portfolioPct"])
            )

    cache.invalidate("gbm_portfolio")


def get_full_portfolio() -> dict:
    """Get complete portfolio data from SQLite with summaries. Cached for 5 minutes."""
    cached = cache.get("gbm_portfolio")
    if cached is not None:
        return cached

    USD_MXN = fetch_usd_mxn()

    with get_db() as conn:
        nac_rows = conn.execute("SELECT * FROM gbm_nacional").fetchall()
        usa_rows = conn.execute("SELECT * FROM gbm_usa").fetchall()

    nacional = []
    for row in nac_rows:
        nacional.append({
            "ticker": row["ticker"],
            "section": row["section"],
            "market": "Nacional",
            "currency": "MXN",
            "shares": row["shares"],
            "avgCost": row["avg_cost"],
            "marketPrice": row["market_price"],
            "marketValue": row["market_value"],
            "gainLoss": row["gain_loss"],
            "returnPct": row["return_pct"],
            "varDayPct": row["var_day_pct"],
            "portfolioPct": row["portfolio_pct"]
        })

    usa = []
    for row in usa_rows:
        market_value = row["market_value"]
        usa.append({
            "ticker": row["ticker"],
            "section": row["section"],
            "market": "USA",
            "currency": "USD",
            "shares": row["shares"],
            "avgCost": row["avg_cost"],
            "marketPrice": row["market_price"],
            "marketValue": market_value,
            "marketValueMXN": round(market_value * USD_MXN, 2),
            "gainLoss": row["gain_loss"],
            "returnPct": row["return_pct"],
            "varDayPct": row["var_day_pct"],
            "costValue": row["cost_value"],
            "portfolioPct": row["portfolio_pct"]
        })

    # Calculate totals
    nacional_value = sum(i["marketValue"] for i in nacional)
    usa_value_usd = sum(i["marketValue"] for i in usa)
    usa_value_mxn = usa_value_usd * USD_MXN
    total_mxn = nacional_value + usa_value_mxn

    nacional_cost = sum(i["avgCost"] * i["shares"] for i in nacional)
    usa_cost = sum(i["costValue"] for i in usa if i.get("costValue", 0) > 0)
    if usa_cost == 0:
        usa_cost = sum(i["avgCost"] * i["shares"] for i in usa)

    total_cost_mxn = nacional_cost + (usa_cost * USD_MXN)
    total_gain = total_mxn - total_cost_mxn
    total_return_pct = (total_gain / total_cost_mxn * 100) if total_cost_mxn > 0 else 0

    result = {
        "nacional": nacional,
        "usa": usa,
        "summary": {
            "nacionalValueMXN": round(nacional_value, 2),
            "usaValueUSD": round(usa_value_usd, 2),
            "usaValueMXN": round(usa_value_mxn, 2),
            "totalValueMXN": round(total_mxn, 2),
            "totalGainMXN": round(total_gain, 2),
            "totalReturnPct": round(total_return_pct, 2),
            "usdMxnRate": USD_MXN
        }
    }

    cache.set("gbm_portfolio", result, ttl_seconds=300)
    return result
