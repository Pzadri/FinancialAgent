"""Migration script: reads all Excel files and populates SQLite database.

Run this once to migrate from Excel to SQLite.
Usage: python migrate_to_sqlite.py
"""
import sys
import json
from pathlib import Path

# Add parent to path so we can import app modules
sys.path.insert(0, str(Path(__file__).parent))

import openpyxl

DATA_DIR = Path(__file__).parent.parent / "data"

from app.database import init_db, get_db


def migrate_ahorro():
    """Migrate ahorro.xlsx to SQLite."""
    filepath = DATA_DIR / "inversiones" / "ahorro.xlsx"
    if not filepath.exists():
        print("  [SKIP] ahorro.xlsx not found")
        return

    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active
    count = 0

    with get_db() as conn:
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            name = str(row[0])
            description = str(row[1]) if row[1] else ""
            color = str(row[2]) if row[2] else "#1da1f2"
            balance = float(row[3]) if row[3] else 0
            annual_rate = float(row[4]) if row[4] else 0

            conn.execute(
                "INSERT INTO ahorro (name, description, color, balance, annual_rate) VALUES (?, ?, ?, ?, ?)",
                (name, description, color, balance, annual_rate)
            )
            count += 1

    wb.close()
    print(f"  [OK] ahorro: {count} records migrated")


def migrate_prestamos():
    """Migrate prestamos.xlsx to SQLite."""
    filepath = DATA_DIR / "inversiones" / "prestamos.xlsx"
    if not filepath.exists():
        print("  [SKIP] prestamos.xlsx not found")
        return

    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active
    count = 0

    with get_db() as conn:
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            borrower = str(row[0])
            principal = float(row[1]) if row[1] else 0
            rate = float(row[2]) if row[2] else 0
            term = str(row[3]) if row[3] else ""
            status = str(row[4]).strip().lower() if row[4] else "activo"

            conn.execute(
                "INSERT INTO prestamos (borrower, principal, rate, term, status) VALUES (?, ?, ?, ?, ?)",
                (borrower, principal, rate, term, status)
            )
            count += 1

    wb.close()
    print(f"  [OK] prestamos: {count} records migrated")


def migrate_afore():
    """Migrate afore.xlsx to SQLite."""
    filepath = DATA_DIR / "inversiones" / "afore.xlsx"
    if not filepath.exists():
        print("  [SKIP] afore.xlsx not found")
        return

    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active

    data = {"balance": 0, "annual_return": 0, "bimonthly_contribution": 0, "voluntary_contribution": 0}

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        field = str(row[0]).strip().lower()
        value = float(row[1]) if row[1] else 0

        if "saldo" in field:
            data["balance"] = value
        elif "rendimiento" in field:
            data["annual_return"] = value
        elif "bimestral" in field:
            data["bimonthly_contribution"] = value
        elif "voluntaria" in field or "semanal" in field:
            data["voluntary_contribution"] = value

    wb.close()

    with get_db() as conn:
        conn.execute(
            "INSERT INTO afore (id, balance, annual_return, bimonthly_contribution, voluntary_contribution) VALUES (1, ?, ?, ?, ?)",
            (data["balance"], data["annual_return"], data["bimonthly_contribution"], data["voluntary_contribution"])
        )

    print(f"  [OK] afore: balance={data['balance']}, return={data['annual_return']}%")


def migrate_creditos():
    """Migrate tarjetas.xlsx to SQLite."""
    filepath = DATA_DIR / "creditos" / "tarjetas.xlsx"
    if not filepath.exists():
        print("  [SKIP] tarjetas.xlsx not found")
        return

    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active
    count = 0

    with get_db() as conn:
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            name = str(row[0])
            color = str(row[1]) if row[1] else "#1da1f2"
            credit_limit = float(row[2]) if row[2] else 0
            debt = float(row[3]) if row[3] else 0
            available = float(row[4]) if row[4] else 0
            cutoff_date = str(row[5]) if row[5] else ""
            payment_date = str(row[6]) if row[6] else ""
            minimum_payment = float(row[7]) if row[7] else 0
            full_payment = float(row[8]) if row[8] else 0

            conn.execute(
                """INSERT INTO creditos (name, color, credit_limit, debt, available, cutoff_date, payment_date, minimum_payment, full_payment)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (name, color, credit_limit, debt, available, cutoff_date, payment_date, minimum_payment, full_payment)
            )
            count += 1

    wb.close()
    print(f"  [OK] creditos: {count} cards migrated")


def migrate_gastos_ingresos():
    """Migrate gastos-ingresos.xlsx to SQLite."""
    filepath = DATA_DIR / "GI" / "gastos-ingresos.xlsx"
    if not filepath.exists():
        print("  [SKIP] gastos-ingresos.xlsx not found")
        return

    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active
    count = 0

    with get_db() as conn:
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            # row[0] is the old ID, we let SQLite auto-increment
            fecha = str(row[1]) if row[1] else ""
            descripcion = str(row[2]) if row[2] else ""
            categoria = str(row[3]) if row[3] else ""
            tipo = str(row[4]) if row[4] else ""
            monto = float(row[5]) if row[5] else 0

            conn.execute(
                "INSERT INTO gastos_ingresos (fecha, descripcion, categoria, tipo, monto) VALUES (?, ?, ?, ?, ?)",
                (fecha, descripcion, categoria, tipo, monto)
            )
            count += 1

    wb.close()
    print(f"  [OK] gastos_ingresos: {count} records migrated")


def migrate_deudas():
    """Migrate deudas.xlsx to SQLite."""
    filepath = DATA_DIR / "deudas" / "deudas.xlsx"
    if not filepath.exists():
        print("  [SKIP] deudas.xlsx not found")
        return

    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active
    count = 0

    with get_db() as conn:
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            nombre = str(row[1]) if row[1] else ""
            deuda_total = float(row[2]) if row[2] else 0
            pago_periodo = float(row[3]) if row[3] else 0
            temporalidad = str(row[4]) if row[4] else "mensual"
            dia1 = int(row[5]) if row[5] else 1
            dia2 = int(row[6]) if row[6] else 0
            pagos_restantes = int(row[7]) if row[7] else 0
            fecha_inicio_1 = str(row[8]) if len(row) > 8 and row[8] else ""
            fecha_inicio_2 = str(row[9]) if len(row) > 9 and row[9] else ""

            conn.execute(
                """INSERT INTO deudas (nombre, deuda_total, pago_periodo, temporalidad, dia_pago_1, dia_pago_2, pagos_restantes, fecha_inicio_1, fecha_inicio_2)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (nombre, deuda_total, pago_periodo, temporalidad, dia1, dia2, pagos_restantes, fecha_inicio_1, fecha_inicio_2)
            )
            count += 1

    wb.close()
    print(f"  [OK] deudas: {count} records migrated")


def migrate_gbm():
    """Migrate GBM Excel files to SQLite."""
    from app.gbm_reader import read_nacional, read_usa, fetch_usd_mxn

    nacional = read_nacional()
    usa = read_usa(fetch_usd_mxn())

    with get_db() as conn:
        for inv in nacional:
            conn.execute(
                """INSERT INTO gbm_nacional (ticker, section, shares, avg_cost, market_price, market_value, gain_loss, return_pct, var_day_pct, portfolio_pct)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (inv["ticker"], inv["section"], inv["shares"], inv["avgCost"], inv["marketPrice"],
                 inv["marketValue"], inv["gainLoss"], inv["returnPct"], inv["varDayPct"], inv["portfolioPct"])
            )

        for inv in usa:
            conn.execute(
                """INSERT INTO gbm_usa (ticker, section, shares, avg_cost, market_price, market_value, gain_loss, return_pct, var_day_pct, cost_value, portfolio_pct)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (inv["ticker"], inv["section"], inv["shares"], inv["avgCost"], inv["marketPrice"],
                 inv["marketValue"], inv["gainLoss"], inv["returnPct"], inv["varDayPct"],
                 inv.get("costValue", 0), inv["portfolioPct"])
            )

    print(f"  [OK] gbm_nacional: {len(nacional)} instruments migrated")
    print(f"  [OK] gbm_usa: {len(usa)} instruments migrated")


def migrate_update_tracker():
    """Migrate last_updates.json to SQLite."""
    tracker_file = DATA_DIR / "last_updates.json"
    if not tracker_file.exists():
        print("  [SKIP] last_updates.json not found")
        return

    data = json.loads(tracker_file.read_text(encoding="utf-8"))

    with get_db() as conn:
        for section, date_str in data.items():
            conn.execute(
                "INSERT OR REPLACE INTO update_tracker (section, last_update) VALUES (?, ?)",
                (section, date_str)
            )

    print(f"  [OK] update_tracker: {len(data)} entries migrated")


def main():
    print("=" * 50)
    print("Financial Dashboard: Excel -> SQLite Migration")
    print("=" * 50)
    print()

    # Initialize schema
    print("[1/9] Initializing database schema...")
    init_db()
    print("  [OK] Schema created")
    print()

    # Migrate each section
    print("[2/9] Migrating ahorro...")
    migrate_ahorro()

    print("[3/9] Migrating prestamos...")
    migrate_prestamos()

    print("[4/9] Migrating afore...")
    migrate_afore()

    print("[5/9] Migrating creditos...")
    migrate_creditos()

    print("[6/9] Migrating gastos/ingresos...")
    migrate_gastos_ingresos()

    print("[7/9] Migrating deudas...")
    migrate_deudas()

    print("[8/9] Migrating GBM portfolio...")
    migrate_gbm()

    print("[9/9] Migrating update tracker...")
    migrate_update_tracker()

    print()
    print("=" * 50)
    print("Migration complete!")
    print(f"Database: {DATA_DIR / 'financial.db'}")
    print("=" * 50)


if __name__ == "__main__":
    main()
